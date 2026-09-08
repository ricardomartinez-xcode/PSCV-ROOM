#!/usr/bin/env python3
import json, os, re, tempfile, subprocess, unicodedata, hashlib
from pathlib import Path
from collections import Counter

raw=json.load(open('r2-audit-materials.json',encoding='utf-8'))
rows=(raw[0].get('results',[]) if isinstance(raw,list) else raw.get('result',[{}])[0].get('results',[]))

def norm(s):
    s=unicodedata.normalize('NFD',s or '')
    return re.sub(r'\s+',' ',''.join(c for c in s if unicodedata.category(c)!='Mn').lower())

def fetch(key):
    with tempfile.NamedTemporaryFile(delete=False) as f:
        path=f.name
    try:
        proc=subprocess.run([
            'npx','wrangler','r2','object','get',f'psicologia/{key}',
            '--file',path,'--remote'
        ],stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=180)
        if proc.returncode != 0:
            raise RuntimeError((proc.stderr or proc.stdout or 'wrangler r2 get failed').strip())
        return Path(path).read_bytes()
    finally:
        try: os.unlink(path)
        except: pass

def extract(data,name):
    ext=Path(name).suffix.lower(); p=None
    try:
        with tempfile.NamedTemporaryFile(suffix=ext,delete=False) as f: f.write(data); p=f.name
        if ext=='.pdf':
            from pypdf import PdfReader
            return ' '.join((pg.extract_text() or '') for pg in PdfReader(p).pages[:100]),'PDF',None
        if ext=='.docx':
            from docx import Document
            d=Document(p); return ' '.join(x.text for x in d.paragraphs),'Documento Word',None
        if ext=='.pptx':
            from pptx import Presentation
            prs=Presentation(p); return ' '.join(sh.text for sl in prs.slides for sh in sl.shapes if hasattr(sh,'text')),'Presentación',None
        if ext=='.xlsx':
            from openpyxl import load_workbook
            wb=load_workbook(p,read_only=True,data_only=True)
            return ' '.join(str(v) for ws in wb.worksheets for row in ws.iter_rows(max_row=500,values_only=True) for v in row if v is not None),'Hoja de cálculo',None
        if ext in {'.txt','.md','.csv','.rtf'}: return data.decode('utf-8','ignore'),'Texto',None
        return '',ext.lstrip('.').upper() or 'Archivo','formato sin extractor'
    except Exception as e: return '',ext.lstrip('.').upper() or 'Archivo',f'{type(e).__name__}: {e}'
    finally:
        if p:
            try: os.unlink(p)
            except: pass

AREAS={
'Psicología Clínica':['psicoterapia','terapia','intervencion','duelo','depresion','ansiedad','paciente','salud mental','diagnostico clinico'],
'Psicopatología':['psicopatologia','trastorno','dsm','cie 10','psicosis','esquizofrenia','bipolar','trastornos de la personalidad'],
'Psicología Educativa':['psicologia educativa','educacion','aprendizaje','docente','alumno','escolar','ensenanza','pedagog'],
'Psicología Organizacional':['psicologia organizacional','recursos humanos','capital humano','talento humano','clima laboral','seleccion de personal'],
'Psicología Jurídica y Forense':['forense','juridica','criminologia','criminal','delito','victima','peritaje','perfilacion'],
'Investigación y Metodología':['metodologia','investigacion','hipotesis','variable','muestra','apa','estadistica','diseno de investigacion'],
'Psicometría':['psicometria','test psicologico','cuestionario','escala','baremo','validez','confiabilidad','wais','wisc','inventario','instrumento psicologico'],
'Psicología Social':['psicologia social','procesos grupales','dinamica de grupos','influencia social','actitud','comunidad','problematica social'],
'Psicología del Desarrollo':['desarrollo humano','desarrollo infantil','adolescencia','infancia','vejez','ciclo vital','piaget','erikson'],
'Neuropsicología':['neuropsicologia','cerebro','funciones ejecutivas','memoria','atencion','neurolog','cognicion'],
'Psicología General':['psicologia','conducta','emocion','motivacion','personalidad','teoria psicologica']}
COURSES={
'Alteraciones de la Conducta':['alteraciones de la conducta','conducta anormal','psicopatologia','trastorno'],
'Evaluación Psicológica I':['evaluacion psicologica','psicodiagnostico','entrevista psicologica','pruebas psicologicas','test psicologico'],
'Psicología del Aprendizaje':['teorias del aprendizaje','psicologia del aprendizaje','condicionamiento','skinner','pavlov','bandura'],
'Teoría y Práctica de Procesos Grupales':['procesos grupales','dinamica de grupos','cohesion grupal','liderazgo grupal'],
'Psicología en la Problemática Social Mexicana':['problematica social mexicana','violencia social','desigualdad','comunidad mexicana']}
TYPES={
'Instrumentos':['test','cuestionario','escala','inventario','protocolo','hoja de respuesta','cuadernillo'],
'Manuales':['manual','guia de aplicacion','instrucciones'],
'Artículos científicos':['abstract','metodologia','resultados','discusion','doi'],
'Casos y prácticas':['caso clinico','caso practico','actividad','ejercicio','practica'],
'Lecturas':['capitulo','lectura','introduccion','conclusion']}

def score(t,words): return sum((3 if ' ' in w else 1)*t.count(norm(w)) for w in words)
def choose(t,table,default):
    ranked=sorted(((score(t,v),k) for k,v in table.items()),reverse=True)
    return (ranked[0][1],ranked[0][0],ranked[1][0] if len(ranked)>1 else 0) if ranked and ranked[0][0] else (default,0,0)
def clean_name(k): return re.sub(r'^[0-9a-f]{8}-[0-9a-f-]{27}-','',Path(k).name,flags=re.I)

COURSE_PREFIXES=[
 ('Alteraciones de la conducta/','Alteraciones de la Conducta'),
 ('Evaluacion Psicológica I/','Evaluación Psicológica I'),
 ('Evaluación Psicológica I/','Evaluación Psicológica I'),
 ('Procesos Grupales/','Teoría y Práctica de Procesos Grupales'),
 ('Teorias del Aprendizaje/','Psicología del Aprendizaje'),
 ('Teorías del Aprendizaje/','Psicología del Aprendizaje'),
 ('Materiales de clase/Sexto cuatrimestre/Alteraciones de la Conducta/','Alteraciones de la Conducta'),
 ('Materiales de clase/Sexto cuatrimestre/Evaluación Psicológica I/','Evaluación Psicológica I'),
 ('Materiales de clase/Sexto cuatrimestre/Teoría y Práctica de Procesos Grupales/','Teoría y Práctica de Procesos Grupales'),
 ('Materiales de clase/Sexto cuatrimestre/Psicología del Aprendizaje/','Psicología del Aprendizaje'),
 ('Materiales de clase/Sexto cuatrimestre/Psicología en la Problemática Social Mexicana/','Psicología en la Problemática Social Mexicana'),
]
INSTRUMENT_DOMAINS={
 'Inteligencia y cognición':['wais','wisc','inteligencia','cognitiv','memoria','herrmann'],
 'Ansiedad y estrés':['ansiedad','stai','estres','estrés','afrontamiento','miedo'],
 'Depresión y estado de ánimo':['depresion','depresión','hamilton','estado de animo','estado de ánimo'],
 'Personalidad':['personalidad','narcis','psicopat','perfil de personalidad'],
 'Adaptación y conducta':['adaptacion','adaptación','procrastin','conducta','cap ado'],
 'Autoestima y autoconcepto':['autoestima','rosenberg','autoconcepto'],
 'Infancia y adolescencia':['infantil','niño','nino','adolescente','menores'],
 'Técnicas proyectivas':['persona bajo la lluvia','proyectiv','dibujo','htp'],
}
def source_course(key):
    for prefix,course in COURSE_PREFIXES:
        if key.startswith(prefix): return course
    return None
def instrument_domain(text):
    nt=norm(text)
    ranked=sorted(((score(nt,v),k) for k,v in INSTRUMENT_DOMAINS.items()),reverse=True)
    return ranked[0][1] if ranked and ranked[0][0]>0 else 'Otros instrumentos'

RESEARCH_TITLE_SIGNALS=['metodologia','metodología','investigacion','investigación','estadistica','estadística','metodo cientifico','método científico']
REFERENCE_TITLE_SIGNALS=['apa 7','codigo etico','código ético','normatividad','lineamientos','guia clinica','guía clínica']
BOOK_TITLE_SIGNALS=['libro','edicion','edición','enciclopedia','tratado','fundamentos de','teorias de la','teorías de la']
INSTRUMENT_TITLE_SIGNALS=['test ','test-','escala','cuestionario','inventario','protocolo','wais','wisc','stai','hamilton','rosenberg','cap ado','nacad','herrmann','persona bajo la lluvia','cuadernillo','hoja de respuesta','plantilla']
ARTICLE_CUES=['doi','abstract','resumen','palabras clave','keywords','resultados','results','discusion','discusión','references','referencias']
STRONG_AREA_TITLE_SIGNALS={
 'Psicología Jurídica y Forense':['criminolog','criminal','forense','juridic','asesino','delito','victima','víctima','perfilacion','perfilación'],
 'Psicopatología':['psicopatolog','trastorno','narcisismo','psicopata','psicópata','psicosis'],
 'Psicología Organizacional':['organizacional','capital humano','recursos humanos','clima laboral','talento humano'],
 'Psicología Educativa':['psicologia educativa','psicología educativa','docente','ensenanza','enseñanza','aprendizaje escolar'],
 'Psicología Clínica':['psicoterapia','intervencion clinica','intervención clínica','duelo','casos clinicos','casos clínicos'],
}

def choose_area(body,title,key):
    t=norm(title+' '+key)
    for area,signals in STRONG_AREA_TITLE_SIGNALS.items():
        if any(norm(x) in t for x in signals):
            return area,12,0
    if any(norm(x) in t for x in RESEARCH_TITLE_SIGNALS):
        return 'Investigación y Metodología', max(8,score(t,AREAS['Investigación y Metodología'])),0
    reduced={k:v for k,v in AREAS.items() if k!='Investigación y Metodología'}
    return choose(body if len(body)>=120 else t,reduced,'Psicología General')

def choose_document_type(title,text,kind):
    tn=norm(title); sample=norm(text[:250000])
    if kind=='Presentación': return 'Presentaciones'
    if any(norm(x) in tn for x in REFERENCE_TITLE_SIGNALS): return 'Guías y normatividad'
    if 'manual' in tn or 'guia de aplicacion' in tn or 'guía de aplicación' in tn: return 'Manuales'
    if any(norm(x) in tn for x in BOOK_TITLE_SIGNALS) or len(text)>180000: return 'Libros y capítulos'
    if any(norm(x) in tn for x in INSTRUMENT_TITLE_SIGNALS): return 'Instrumentos'
    if len(text)>=1500 and len(text)<=180000 and sum(1 for x in ARTICLE_CUES if norm(x) in sample)>=4: return 'Artículos científicos'
    if any(x in tn for x in ['articulo','artículo','boletin','boletín']) and len(text)<=220000: return 'Artículos científicos'
    if any(x in tn for x in ['caso','actividad','ejercicio','practica','práctica','tarea','proyecto final']): return 'Casos y prácticas'
    if any(x in sample for x in ['caso clinico','caso clínico','caso practico','caso práctico']) and len(text)<120000: return 'Casos y prácticas'
    return 'Lecturas'

def is_instrument_resource(title,dtype,key):
    tn=norm(title); kn=norm(key)
    return dtype=='Instrumentos' or (
        dtype=='Manuales' and (
            any(norm(x) in tn for x in INSTRUMENT_TITLE_SIGNALS)
            or 'test cuestionarios' in kn
            or 'cuestionario de adaptacion' in kn
        )
    )

out=[]
for i,row in enumerate(rows,1):
    key=row['r2_key']; title=row.get('title') or clean_name(key)
    try:
        data=fetch(key); sha256=hashlib.sha256(data).hexdigest(); byte_size=len(data); text,kind,err=extract(data,key)
    except Exception as e:
        data=b''; sha256=None; byte_size=0; text=''; kind='Archivo'; err=f'{type(e).__name__}: {e}'
    body=norm(text[:800000]); basis=body if len(body)>=120 else norm(title+' '+key)
    area,ascore,a2=choose_area(body,title,key)
    inferred_course,cscore,c2=choose(basis,COURSES,'')
    course=source_course(key)
    dtype=choose_document_type(title,text,kind); dscore=0
    title_norm=norm(title)
    psychometric_signal=is_instrument_resource(title,dtype,key)
    if course:
        target=f'Materias/Sexto cuatrimestre/{course}/{dtype}/{clean_name(key)}'; section=f'Materias / Sexto cuatrimestre / {course} / {dtype}'
    elif psychometric_signal:
        domain=instrument_domain((title+' ')*8+(key+' ')*3+body[:60000])
        sub='Manuales' if ('manual' in title_norm or dtype=='Manuales') else 'Instrumentos'
        target=f'Instrumentos psicológicos/{domain}/{sub}/{clean_name(key)}'; section=f'Instrumentos psicológicos / {domain} / {sub}'
    else:
        target=f'Biblioteca/{area}/{dtype}/{clean_name(key)}'; section=f'Biblioteca / {area} / {dtype}'
    strength=max(ascore,cscore); conf='alta' if strength>=12 else ('media' if strength>=5 else 'baja')
    rec={'id':row.get('id'),'title':title,'source':key,'target':target,'section':section,'area':area,'course':course,'inferred_course':inferred_course or None,'document_type':dtype,'file_kind':kind,'text_chars':len(text),'byte_size':byte_size,'sha256':sha256,'confidence':conf,'scores':{'area':ascore,'course':cscore,'type':dscore},'needs_review':bool(err or len(text)<120 or (conf=='baja' and not course))}
    if err: rec['extraction_error']=err
    out.append(rec); print(f'[{i}/{len(rows)}] {conf:5} {area} :: {title[:90]} :: error={err or "-"}')

Path('outputs').mkdir(exist_ok=True)
json.dump(out,open('outputs/r2-semantic-classification.json','w',encoding='utf-8'),ensure_ascii=False,indent=2)
hash_groups={}
for x in out:
    if x.get('sha256'): hash_groups.setdefault(x['sha256'],[]).append(x)
duplicate_groups=[g for g in hash_groups.values() if len(g)>1]
summary={'total':len(out),'unique_hashes':len(hash_groups),'exact_duplicate_groups':len(duplicate_groups),'exact_duplicate_records':sum(len(g) for g in duplicate_groups),'readable_content':sum(x['text_chars']>=120 for x in out),'needs_review':sum(x['needs_review'] for x in out),'confidence':dict(Counter(x['confidence'] for x in out)),'areas':dict(Counter(x['area'] for x in out)),'document_types':dict(Counter(x['document_type'] for x in out)),'courses':dict(Counter(x['course'] for x in out if x['course'])),'sections':dict(Counter(x['section'] for x in out)),'extraction_errors':sum('extraction_error' in x for x in out),'duplicate_groups':[{'sha256':g[0]['sha256'],'count':len(g),'title':g[0]['title'],'sources':[x['source'] for x in g]} for g in duplicate_groups],'review_items':[{'title':x['title'],'source':x['source'],'area':x['area'],'course':x['course'],'text_chars':x['text_chars'],'extraction_error':x.get('extraction_error')} for x in out if x['needs_review']]}
plan=[]
for g in duplicate_groups:
    targets=sorted(set(x['target'] for x in g))
    preferred=sorted(g,key=lambda x:(not bool(x.get('course')),x.get('needs_review',True),0 if x['source'].startswith('Materiales de clase/') else 1,x['source']))[0]
    safe=(len(targets)==1 and (bool(preferred.get('course')) or not preferred.get('needs_review')))
    plan.append({'sha256':preferred['sha256'],'canonical_source':preferred['source'],'target':preferred['target'],'duplicate_sources':[x['source'] for x in g if x['source']!=preferred['source']],'safe_to_migrate':safe,'needs_review':not safe,'target_conflicts':targets if len(targets)>1 else []})
summary['migration_plan']={'groups':len(plan),'safe_groups':sum(x['safe_to_migrate'] for x in plan),'review_groups':sum(x['needs_review'] for x in plan)}
json.dump(plan,open('outputs/r2-semantic-migration-plan.json','w',encoding='utf-8'),ensure_ascii=False,indent=2)
json.dump(summary,open('outputs/r2-semantic-summary.json','w',encoding='utf-8'),ensure_ascii=False,indent=2)
print(json.dumps(summary,ensure_ascii=False,indent=2))
